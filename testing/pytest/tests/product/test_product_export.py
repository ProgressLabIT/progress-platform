"""Integration tests for GET /product/export.

These tests require a running ArangoDB testcontainer (no Docker mock).
They do NOT seed any product data — header row, template-only, and CSV header
assertions hold on an empty catalog.
"""
import csv as csv_mod
import io

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.asyncio

EXPECTED_COLUMNS = [
    "code", "description", "tags", "print_templates", "counter",
    "manage_inventory", "default_consumption_position", "default_production_position",
]


async def test_product_export_xlsx_headers(client, auth_headers):
    """GET /product/export returns xlsx with the 8-column contract header row (RTR-02)."""
    resp = await client.get("/product/export", params={"format": "xlsx"}, headers=auth_headers())
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert [ws.cell(1, c + 1).value for c in range(8)] == EXPECTED_COLUMNS


async def test_product_export_template_is_header_only(client, auth_headers):
    """format=template returns an xlsx with exactly the header row, no data rows (EXP-04)."""
    resp = await client.get("/product/export", params={"format": "template"}, headers=auth_headers())
    assert resp.status_code == 200
    assert load_workbook(io.BytesIO(resp.content)).active.max_row == 1


async def test_product_export_csv_headers(client, auth_headers):
    """format=csv returns CSV with the same 8-column header row (EXP-02)."""
    resp = await client.get("/product/export", params={"format": "csv"}, headers=auth_headers())
    assert resp.status_code == 200
    assert "text/csv" in resp.headers["content-type"]
    reader = csv_mod.reader(io.StringIO(resp.content.decode("utf-8-sig")))
    assert next(reader) == EXPECTED_COLUMNS
