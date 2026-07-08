"""Integration tests for POST /product/import.

Tests require a running ArangoDB testcontainer (no Docker mock).
Master data (Tag, Counter, PrintTemplate, Position) is seeded via direct collection
inserts — these are reference entities with no cascade side-effects.
Product entities that need to exist before import are seeded via direct collection
inserts; Product has no cascade side-effects for the fields we test (RESEARCH Pitfall 4).

All tests follow the two-call pattern:
  1. POST dry_run=True  (file upload) → file_key
  2. POST dry_run=False (file_key)    → created products
"""

import io
import json
import os

import pytest
from openpyxl import Workbook, load_workbook

pytestmark = pytest.mark.asyncio

EXPORT_COLUMNS = [
    'code', 'description', 'tags', 'print_templates', 'counter',
    'manage_inventory', 'default_consumption_position', 'default_production_position',
]

# ============================================================
# Helpers
# ============================================================

def _make_xlsx(rows: list[dict], columns: list[str] = EXPORT_COLUMNS) -> bytes:
    """Build an in-memory xlsx with the given column headers and data rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Products'
    ws.append(columns)
    for row in rows:
        ws.append([row.get(c, '') for c in columns])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _make_csv(rows: list[dict], columns: list[str] = EXPORT_COLUMNS) -> bytes:
    """Build an in-memory CSV with the given column headers and data rows."""
    import csv as csv_mod
    buf = io.StringIO()
    writer = csv_mod.DictWriter(buf, fieldnames=columns, extrasaction='ignore')
    writer.writeheader()
    writer.writerows(rows)
    # BOM-prefix for Excel compat (matches export endpoint convention)
    return ('﻿' + buf.getvalue()).encode('utf-8')


async def _dry_run(client, auth_headers, content: bytes, filename: str = 'test.xlsx'):
    """POST dry_run=True and return the response."""
    media_type = (
        'text/csv'
        if filename.endswith('.csv')
        else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    return await client.post(
        '/product/import',
        files={'file': (filename, content, media_type)},
        data={'dry_run': 'true'},
        headers=auth_headers(),
    )


async def _execute(client, auth_headers, file_key: str):
    """POST dry_run=False with the given file_key and return the response."""
    return await client.post(
        '/product/import',
        data={'dry_run': 'false', 'file_key': file_key},
        headers=auth_headers(),
    )


def _seed_product(db, code: str, **overrides) -> str:
    """
    Insert a product directly for test setup (no cascade for Product docs — RESEARCH Pitfall 4).
    Returns the inserted _key.
    """
    doc = {
        'code': code,
        'active': True,
        'trash': False,
        **overrides,
    }
    return db.collection('Product').insert(doc)['_key']


# ============================================================
# Tests — Plan 01 (create-path)
# ============================================================

async def test_import_creates_new_product(client, auth_headers, db):
    """IMP-01/IMP-03: xlsx upload creates product through event pipeline."""
    rows = [{'code': 'TESTIMP-001', 'description': 'Import test product'}]
    content = _make_xlsx(rows)

    # Dry run
    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body['status'] == 'valid'
    assert body['created_count'] == 1
    file_key = body['file_key']

    # Execute
    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text
    result = resp2.json()
    assert result['detail']['created'] == 1

    # Verify product exists in DB
    products = list(db.aql.execute(
        "FOR p IN Product FILTER p.code == @code RETURN p",
        bind_vars={'code': 'TESTIMP-001'}
    ))
    assert len(products) == 1
    assert products[0]['description'] == 'Import test product'


async def test_import_creates_from_csv(client, auth_headers, db):
    """IMP-02: CSV upload creates product with same result as xlsx."""
    rows = [{'code': 'TESTIMP-CSV-001', 'description': 'CSV import test'}]
    content = _make_csv(rows)

    resp = await _dry_run(client, auth_headers, content, filename='test.csv')
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body['status'] == 'valid'
    file_key = body['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text
    assert resp2.json()['detail']['created'] == 1

    products = list(db.aql.execute(
        "FOR p IN Product FILTER p.code == @code RETURN p",
        bind_vars={'code': 'TESTIMP-CSV-001'}
    ))
    assert len(products) == 1
    assert products[0]['description'] == 'CSV import test'


async def test_import_applies_tag_edges(client, auth_headers, db):
    """IMP-05/REF-01: has_tag edge is created from the new Product to the seeded Tag."""
    tag_key = db.collection('Tag').insert({'name': 'steel'})['_key']

    rows = [{'code': 'TESTIMP-TAG-001', 'tags': 'steel'}]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    file_key = resp.json()['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text

    # Find the created product
    products = list(db.aql.execute(
        "FOR p IN Product FILTER p.code == @code RETURN p._key",
        bind_vars={'code': 'TESTIMP-TAG-001'}
    ))
    assert products, "Product was not created"
    product_key = products[0]

    # Verify has_tag edge
    edges = list(db.aql.execute(
        "FOR e IN has_tag FILTER e._from == @prod_id RETURN e",
        bind_vars={'prod_id': f'Product/{product_key}'}
    ))
    assert len(edges) == 1
    assert edges[0]['_to'] == f'Tag/{tag_key}'


async def test_import_applies_print_template_edges(client, auth_headers, db):
    """IMP-05/REF-02: can_use_print_template edge is created from the new Product to the seeded template."""
    tmpl_key = db.collection('PrintTemplate').insert({'name': 'label-a4'})['_key']

    rows = [{'code': 'TESTIMP-TPL-001', 'print_templates': 'label-a4'}]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    file_key = resp.json()['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text

    products = list(db.aql.execute(
        "FOR p IN Product FILTER p.code == @code RETURN p._key",
        bind_vars={'code': 'TESTIMP-TPL-001'}
    ))
    assert products, "Product was not created"
    product_key = products[0]

    edges = list(db.aql.execute(
        "FOR e IN can_use_print_template FILTER e._from == @prod_id RETURN e",
        bind_vars={'prod_id': f'Product/{product_key}'}
    ))
    assert len(edges) == 1
    assert edges[0]['_to'] == f'PrintTemplate/{tmpl_key}'


async def test_import_resolves_counter_and_positions(client, auth_headers, db):
    """REF-03/REF-04: counter_key and both position keys are set on the created product."""
    counter_key = db.collection('Counter').insert({'name': 'SN-TEST', 'next_tick': 1})['_key']
    cons_pos_key = db.collection('Position').insert({'code': 'WH-CONS'})['_key']
    prod_pos_key = db.collection('Position').insert({'code': 'WH-PROD'})['_key']

    rows = [{
        'code': 'TESTIMP-REF-001',
        'counter': 'SN-TEST',
        'default_consumption_position': 'WH-CONS',
        'default_production_position': 'WH-PROD',
    }]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    file_key = resp.json()['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text

    products = list(db.aql.execute(
        "FOR p IN Product FILTER p.code == @code RETURN p",
        bind_vars={'code': 'TESTIMP-REF-001'}
    ))
    assert products, "Product was not created"
    product = products[0]
    assert product['counter_key'] == counter_key
    assert product['default_consumption_position_key'] == cons_pos_key
    assert product['default_production_position_key'] == prod_pos_key


async def test_import_unknown_ref_returns_error_file(client, auth_headers, db):
    """REF-01..04 / D-03: unknown tag name → dry_run returns annotated xlsx, Product count unchanged."""
    initial_count = db.collection('Product').count()

    rows = [{'code': 'TESTIMP-ERR-001', 'tags': 'nonexistent-tag-xyz'}]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text

    # Response should be an error xlsx
    content_type = resp.headers.get('content-type', '')
    assert 'spreadsheetml' in content_type or 'excel' in content_type, (
        f"Expected spreadsheet content-type, got: {content_type}"
    )
    assert int(resp.headers['x-error-count']) >= 1

    # Product count unchanged — nothing committed
    assert db.collection('Product').count() == initial_count


async def test_import_through_event_pipeline_writes_event(client, auth_headers, db):
    """IMP-03: after a successful execute, the Event collection has a PRODUCT_IMPORTED document."""
    event_count_before = db.collection('Event').count()

    rows = [{'code': 'TESTIMP-EVT-001', 'description': 'Event pipeline test'}]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    file_key = resp.json()['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text

    # Confirm a PRODUCT_IMPORTED event was written
    events = list(db.aql.execute(
        "FOR e IN Event FILTER e.event_type == 'PRODUCT_IMPORTED' RETURN e"
    ))
    assert len(events) >= 1, "No PRODUCT_IMPORTED event found in Event collection"

    # Also confirm event count increased (proves Event.save() was called, not a raw insert)
    assert db.collection('Event').count() > event_count_before


async def test_execute_requires_validated_file_key(client, auth_headers):
    """D-03: executing with a bogus file_key returns 4xx."""
    resp = await _execute(client, auth_headers, 'bogus-file-key-does-not-exist')
    assert resp.status_code in (400, 404, 422), (
        f"Expected 4xx for bogus file_key, got {resp.status_code}: {resp.text}"
    )


# ============================================================
# Tests — Plan 02 (upsert / idempotency)
# ============================================================

async def test_import_updates_existing_product_by_code(client, auth_headers, db):
    """IMP-04: existing product matched by code; description updated in place; code unchanged."""
    # Seed an existing product with code X and old description
    product_key = _seed_product(db, code='UPSERT-001', description='Old description')

    # Import a file with the same code and a changed description
    rows = [{'code': 'UPSERT-001', 'description': 'New description'}]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    file_key = resp.json()['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text
    result = resp2.json()['detail']
    assert result['updated'] == 1
    assert result['created'] == 0

    # Exactly one product with code UPSERT-001
    products = list(db.aql.execute(
        "FOR p IN Product FILTER p.code == @code AND !p.trash RETURN p",
        bind_vars={'code': 'UPSERT-001'}
    ))
    assert len(products) == 1, f"Expected 1 product with code UPSERT-001, got {len(products)}"
    assert products[0]['description'] == 'New description'
    # code must never change
    assert products[0]['code'] == 'UPSERT-001'
    # _key must be the same (same document, not a new insert)
    assert products[0]['_key'] == product_key


async def test_import_idempotent_no_op(client, auth_headers, db):
    """IMP-06/D-02: re-importing the unchanged file produces zero writes on the second run."""
    rows = [{'code': 'IDEM-001', 'description': 'Idempotent product'}]
    content = _make_xlsx(rows)

    # First import — creates the product
    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    file_key = resp.json()['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text
    assert resp2.json()['detail']['created'] == 1

    # Capture state after first import
    product_before = list(db.aql.execute(
        "FOR p IN Product FILTER p.code == @code AND !p.trash RETURN p",
        bind_vars={'code': 'IDEM-001'}
    ))[0]
    rev_before = product_before['_rev']

    # Second import — same file content, same file_key cannot be reused; re-upload
    resp3 = await _dry_run(client, auth_headers, content)
    assert resp3.status_code == 200, resp3.text
    file_key2 = resp3.json()['file_key']

    resp4 = await _execute(client, auth_headers, file_key2)
    assert resp4.status_code == 200, resp4.text
    result2 = resp4.json()['detail']

    # Headline assertion: second run created == 0 AND updated == 0
    assert result2['created'] == 0, f"Second import should not create; got {result2}"
    assert result2['updated'] == 0, f"Second import should not update; got {result2}"
    assert result2['skipped'] >= 1

    # The product doc must not have been modified (_rev unchanged)
    product_after = list(db.aql.execute(
        "FOR p IN Product FILTER p.code == @code AND !p.trash RETURN p",
        bind_vars={'code': 'IDEM-001'}
    ))[0]
    assert product_after['_rev'] == rev_before, "Product _rev changed on no-op re-import"


async def test_import_clears_field_when_column_present_empty(client, auth_headers, db):
    """D-01: present-but-empty description → ""; present-but-empty tags → 0 has_tag edges."""
    # Seed a tag and a product that has it
    tag_key = db.collection('Tag').insert({'name': 'tag-to-clear'})['_key']
    product_key = _seed_product(db, code='CLEAR-001', description='old description')
    # Manually insert a has_tag edge
    db.collection('has_tag').insert({
        '_from': f'Product/{product_key}',
        '_to': f'Tag/{tag_key}',
    })

    # Import: description column present-but-empty, tags column present-but-empty
    rows = [{'code': 'CLEAR-001', 'description': '', 'tags': ''}]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    file_key = resp.json()['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text
    result = resp2.json()['detail']
    assert result['updated'] == 1

    # description must be cleared to ""
    products = list(db.aql.execute(
        "FOR p IN Product FILTER p.code == @code AND !p.trash RETURN p",
        bind_vars={'code': 'CLEAR-001'}
    ))
    assert len(products) == 1
    assert products[0].get('description', '') == '', (
        f"Expected description='', got: {products[0].get('description')!r}"
    )

    # has_tag edges must be cleared
    edges = list(db.aql.execute(
        "FOR e IN has_tag FILTER e._from == @prod_id RETURN e",
        bind_vars={'prod_id': f'Product/{product_key}'}
    ))
    assert len(edges) == 0, f"Expected 0 has_tag edges after clear, got {len(edges)}"


async def test_import_leaves_field_when_column_absent(client, auth_headers, db):
    """D-01: counter column absent from file header → existing counter_key left untouched."""
    counter_key = db.collection('Counter').insert({'name': 'SN-KEEP', 'next_tick': 1})['_key']
    # Seed product with counter_key set
    _seed_product(db, code='NOCHANGE-001', counter_key=counter_key)

    # Import file that omits the counter column entirely
    columns_without_counter = [c for c in EXPORT_COLUMNS if c != 'counter']
    rows = [{'code': 'NOCHANGE-001', 'description': 'Updated description'}]
    content = _make_xlsx(rows, columns=columns_without_counter)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    file_key = resp.json()['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text

    # counter_key must still be set to the original value
    products = list(db.aql.execute(
        "FOR p IN Product FILTER p.code == @code AND !p.trash RETURN p",
        bind_vars={'code': 'NOCHANGE-001'}
    ))
    assert len(products) == 1
    assert products[0].get('counter_key') == counter_key, (
        f"counter_key should remain {counter_key!r}, got {products[0].get('counter_key')!r}"
    )


async def test_import_edge_set_diff_is_order_insensitive(client, auth_headers, db):
    """D-02: tags in different order than stored → classified as identical → updated==0."""
    tag_a_key = db.collection('Tag').insert({'name': 'alpha'})['_key']
    tag_b_key = db.collection('Tag').insert({'name': 'beta'})['_key']

    # Seed product with tags [alpha, beta]
    product_key = _seed_product(db, code='ORDER-001', description='order test')
    db.collection('has_tag').insert({'_from': f'Product/{product_key}', '_to': f'Tag/{tag_a_key}'})
    db.collection('has_tag').insert({'_from': f'Product/{product_key}', '_to': f'Tag/{tag_b_key}'})

    # Import with tags "beta;alpha" (reversed order)
    rows = [{'code': 'ORDER-001', 'description': 'order test', 'tags': 'beta;alpha'}]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    file_key = resp.json()['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text
    result = resp2.json()['detail']

    # Reordering alone is not a change — row must be skipped
    assert result['updated'] == 0, f"Reordered tags should not trigger update; got {result}"
    assert result['skipped'] >= 1

    # Edges must remain unchanged (still {alpha, beta})
    edges = list(db.aql.execute(
        "FOR e IN has_tag FILTER e._from == @prod_id RETURN e._to",
        bind_vars={'prod_id': f'Product/{product_key}'}
    ))
    assert set(edges) == {f'Tag/{tag_a_key}', f'Tag/{tag_b_key}'}


async def test_import_changed_edge_set_updates(client, auth_headers, db):
    """IMP-04/D-02: existing tags [A,B]; import tags [B,C] → final set is exactly {B,C}."""
    tag_a_key = db.collection('Tag').insert({'name': 'tagA'})['_key']
    tag_b_key = db.collection('Tag').insert({'name': 'tagB'})['_key']
    tag_c_key = db.collection('Tag').insert({'name': 'tagC'})['_key']

    # Seed product with tags [A, B]
    product_key = _seed_product(db, code='EDGESET-001', description='edge set test')
    db.collection('has_tag').insert({'_from': f'Product/{product_key}', '_to': f'Tag/{tag_a_key}'})
    db.collection('has_tag').insert({'_from': f'Product/{product_key}', '_to': f'Tag/{tag_b_key}'})

    # Import with tags "tagB;tagC" (A removed, C added)
    rows = [{'code': 'EDGESET-001', 'description': 'edge set test', 'tags': 'tagB;tagC'}]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    file_key = resp.json()['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text
    result = resp2.json()['detail']
    assert result['updated'] == 1, f"Expected 1 update; got {result}"

    # Final has_tag edge set must be exactly {B, C}
    edges = list(db.aql.execute(
        "FOR e IN has_tag FILTER e._from == @prod_id RETURN e._to",
        bind_vars={'prod_id': f'Product/{product_key}'}
    ))
    assert set(edges) == {f'Tag/{tag_b_key}', f'Tag/{tag_c_key}'}, (
        f"Expected {{B, C}}, got {set(edges)}"
    )


# ============================================================
# Tests — Plan 03-01 (VAL-01..04)
# ============================================================

async def test_val01_dry_run_per_row_classification(client, auth_headers, db):
    """VAL-01: dry-run rows[] carries per-row create/update classification; first row has row==2."""
    # Pre-seed one product so it shows up as 'update'
    _seed_product(db, code='VAL01-EXISTING', description='pre-existing')

    rows = [
        {'code': 'VAL01-NEW', 'description': 'will be created'},
        {'code': 'VAL01-EXISTING', 'description': 'will be updated'},
    ]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body['status'] == 'valid'
    assert 'rows' in body, "dry-run response must include rows[]"
    assert len(body['rows']) == 2

    # First data row (xlsx row 2) must have row==2
    assert body['rows'][0]['row'] == 2, (
        f"First row entry must have row=2 (xlsx convention), got {body['rows'][0]['row']}"
    )
    assert body['rows'][0]['action'] == 'create', (
        f"New code should be action='create', got {body['rows'][0]['action']}"
    )
    assert body['rows'][1]['action'] == 'update', (
        f"Existing code should be action='update', got {body['rows'][1]['action']}"
    )
    assert body['rows'][0]['code'] == 'VAL01-NEW'
    assert body['rows'][1]['code'] == 'VAL01-EXISTING'


async def test_val02_row_number_agreement_xlsx(client, auth_headers, db):
    """VAL-02: first data row error lands on xlsx row 2 (off-by-one fix confirmed)."""
    # A row with an unresolved tag reference — guaranteed to error
    rows = [{'code': 'VAL02-ERR-001', 'tags': 'nonexistent-tag-xyz-val02'}]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text

    # Response should be the error xlsx (not JSON)
    content_type = resp.headers.get('content-type', '')
    assert 'spreadsheetml' in content_type or 'excel' in content_type, (
        f"Expected spreadsheet for error path, got: {content_type}"
    )
    assert int(resp.headers.get('x-error-count', 0)) >= 1

    # Load the xlsx and verify the FIRST data row (xlsx row 2) has status='ERROR'
    # status column is at position len(EXPORT_COLUMNS) + 1 = 9
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    status_col = len(EXPORT_COLUMNS) + 1  # column 9
    # Row 1 = header; row 2 = first data row
    status_cell_value = ws.cell(row=2, column=status_col).value
    assert status_cell_value == 'ERROR', (
        f"Expected xlsx row 2 status='ERROR' (off-by-one fix), got {status_cell_value!r}. "
        "If this is None or 'OK', the enumerate(rows, start=2) fix is not working."
    )


async def test_val02_type_mismatch_manage_inventory(client, auth_headers, db):
    """VAL-02: manage_inventory='yes' (non-bool) produces a type-mismatch error, not silent False."""
    rows = [{'code': 'VAL02-TM-001', 'manage_inventory': 'yes'}]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text

    # Should return the error xlsx (error path because manage_inventory='yes' is invalid)
    content_type = resp.headers.get('content-type', '')
    assert 'spreadsheetml' in content_type or 'excel' in content_type, (
        f"Expected error xlsx for type-mismatch, got JSON (silent coercion still active?): {resp.text[:200]}"
    )
    assert int(resp.headers.get('x-error-count', 0)) >= 1

    # Load xlsx and verify the error message mentions type-mismatch
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    errors_col = len(EXPORT_COLUMNS) + 2  # 'errors' column follows 'status'
    error_cell_value = ws.cell(row=2, column=errors_col).value or ''
    assert 'type-mismatch' in error_cell_value.lower(), (
        f"Expected 'type-mismatch' in error cell, got: {error_cell_value!r}"
    )


async def test_val03_mixed_file_writes_nothing(client, auth_headers, db):
    """VAL-03: mixed valid/invalid file → dry-run returns error xlsx; no products written (block-all)."""
    rows = [
        {'code': 'VAL03-VALID-001', 'description': 'valid row'},
        {'code': 'VAL03-INVALID-001', 'tags': 'nonexistent-tag-xyz-val03'},
    ]
    content = _make_xlsx(rows)

    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text

    # Error path — xlsx returned
    content_type = resp.headers.get('content-type', '')
    assert 'spreadsheetml' in content_type or 'excel' in content_type, (
        f"Expected error xlsx for mixed file, got: {content_type}"
    )
    assert int(resp.headers.get('x-error-count', 0)) >= 1

    # Block-all: neither product must exist in the database
    codes = ['VAL03-VALID-001', 'VAL03-INVALID-001']
    products = list(db.aql.execute(
        'FOR p IN Product FILTER p.code IN @codes RETURN p',
        bind_vars={'codes': codes}
    ))
    assert len(products) == 0, (
        f"Block-all failed: {len(products)} product(s) written from mixed dry-run: "
        f"{[p['code'] for p in products]}"
    )


async def test_val04_execute_summary_counts(client, auth_headers, db):
    """VAL-04: successful execute returns created=1/updated=0/skipped=0; re-execute gives skipped=1."""
    rows = [{'code': 'VAL04-001', 'description': 'summary count test'}]
    content = _make_xlsx(rows)

    # First import: create
    resp = await _dry_run(client, auth_headers, content)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body['status'] == 'valid'
    file_key = body['file_key']

    resp2 = await _execute(client, auth_headers, file_key)
    assert resp2.status_code == 200, resp2.text
    detail = resp2.json()['detail']
    assert detail['created'] == 1, f"Expected created=1, got {detail}"
    assert detail['updated'] == 0, f"Expected updated=0, got {detail}"
    assert detail['skipped'] == 0, f"Expected skipped=0, got {detail}"

    # Second import (same content, same file — re-upload required after first execute)
    resp3 = await _dry_run(client, auth_headers, content)
    assert resp3.status_code == 200, resp3.text
    file_key2 = resp3.json()['file_key']

    resp4 = await _execute(client, auth_headers, file_key2)
    assert resp4.status_code == 200, resp4.text
    detail2 = resp4.json()['detail']
    assert detail2['skipped'] == 1, (
        f"Expected skipped=1 on idempotent re-execute, got {detail2}"
    )
    assert detail2['created'] == 0, f"Expected created=0 on re-execute, got {detail2}"
